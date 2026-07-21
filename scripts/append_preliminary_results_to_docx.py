"""Append the Markdown source of the verified results section to a DOCX copy."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "LLM_Token_Load_Forecasting_Research_Plan_preliminary_results.docx"
MARKDOWN = ROOT / "docs" / "results_discussion_preliminary.md"
OUTPUT = ROOT / "LLM_Token_Load_Forecasting_Research_Plan_preliminary_results_1500.docx"


def read_markdown() -> tuple[str, list[str]]:
    """Read one top-level heading and body paragraphs from the Markdown source."""
    heading = None
    paragraphs: list[str] = []
    for raw_line in MARKDOWN.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            heading = line[2:]
        elif not line.startswith("#"):
            paragraphs.append(line.replace("`", ""))
    if heading is None:
        raise ValueError(f"No top-level heading found in {MARKDOWN}")
    return heading, paragraphs


def remove_prior_results_section(document: Document) -> bool:
    """Remove the short preliminary section appended in the prior document run."""
    remove_from_here = False
    removed = False
    for paragraph in list(document.paragraphs):
        if paragraph.text.startswith("7. Results and Discussion"):
            remove_from_here = True
        if remove_from_here:
            paragraph._element.getparent().remove(paragraph._element)
            removed = True
    return removed


def main() -> None:
    heading, paragraphs = read_markdown()
    document = Document(SOURCE)
    had_prior_section = remove_prior_results_section(document)
    if not had_prior_section:
        page_break = document.add_paragraph()
        page_break.add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph(heading, style="Heading 1")
    for text in paragraphs:
        paragraph = document.add_paragraph(text, style="Normal")
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.3
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
