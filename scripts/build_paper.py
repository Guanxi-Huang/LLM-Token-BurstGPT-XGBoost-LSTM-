"""Build the traceable Word/PDF-ready paper draft from frozen CSV/PNG outputs."""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper" / "paper_draft.md"
TABLES = ROOT / "artifacts" / "final" / "tables"
FIGURES = ROOT / "artifacts" / "final" / "figures"
DELIVERABLES = ROOT / "deliverables"
OUTPUT = DELIVERABLES / "LLM_Token_Load_Forecasting_Draft.docx"

NAVY = "17324D"
BLUE = "2F6B8A"
PALE_BLUE = "EAF2F7"
PALE_GRAY = "F4F6F8"
MID_GRAY = "66717C"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_no_wrap(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    if tc_pr.find(qn("w:noWrap")) is None:
        tc_pr.append(OxmlElement("w:noWrap"))


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=None, color=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline_markup(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, "Consolas", 9.5, color=NAVY)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True, color=NAVY)
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def add_caption(document: Document, text: str, kind: str) -> None:
    paragraph = document.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, size=9, bold=True, color=NAVY)
    paragraph.paragraph_format.keep_with_next = kind == "table"


def add_note(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.left_indent = Cm(0.2)
    run = paragraph.add_run("注：" + text)
    set_run_font(run, size=8.5, color=MID_GRAY)


def add_table(
    document: Document,
    caption: str,
    headers: list[str],
    rows: list[list[str]],
    note: str,
    widths_cm: list[float] | None = None,
) -> None:
    add_caption(document, caption, "table")
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = False
    table.style = "Table Grid"
    prevent_row_split(table.rows[0])
    repeat_header_row(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if widths_cm:
            cell.width = Cm(widths_cm[index])
        if len(header) <= 12 and " " not in header:
            set_cell_no_wrap(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, size=8.5, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths_cm:
                cell.width = Cm(widths_cm[col_index])
            if len(str(value)) <= 12 and " " not in str(value):
                set_cell_no_wrap(cell)
            if row_index % 2:
                set_cell_shading(cell, PALE_GRAY)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(run, size=8.3)
        prevent_row_split(table.rows[-1])
    if widths_cm:
        for index, width in enumerate(widths_cm):
            for cell in table.columns[index].cells:
                cell.width = Cm(width)
    add_note(document, note)


def add_figure(document: Document, path: Path, caption: str, width_inches=6.45) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width_inches))
    add_caption(document, caption, "figure")


def exact_main_rows() -> list[list[str]]:
    baseline = pd.read_csv(TABLES / "table_02_baseline_results.csv")
    baseline = baseline[(baseline["split"] == "test") & (baseline["horizon_minutes"] == 15)]
    xgb = pd.read_csv(TABLES / "table_04_xgb_results.csv")
    xgb = xgb[(xgb["split"] == "test") & (xgb["horizon_minutes"] == 15)].iloc[0]
    lstm = pd.read_csv(TABLES / "table_05_lstm_results.csv")
    lstm = lstm[(lstm["split"] == "test") & (lstm["horizon_minutes"] == 15)].iloc[0]
    burst = pd.read_csv(TABLES / "table_05_burst_results.csv").set_index("model")
    source = {
        "Persistence": baseline[baseline["model"] == "persistence"].iloc[0],
        "Seasonal Naive": baseline[baseline["model"] == "seasonal_naive"].iloc[0],
        "XGBoost": xgb,
        "LSTM": lstm,
    }
    seasonal_mae = float(source["Seasonal Naive"]["mae"])
    rows = []
    for model in ("Persistence", "Seasonal Naive", "XGBoost", "LSTM"):
        row = source[model]
        mae = float(row["mae"])
        rows.append(
            [
                model,
                f"{mae:,.1f}",
                f"{float(row['rmse']):,.1f}",
                f"{float(row['smape_percent']):.1f}%",
                f"{100 * (seasonal_mae - mae) / seasonal_mae:.1f}%",
                f"{float(burst.loc[model, 'pr_auc']):.4f}",
            ]
        )
    return rows


def insert_data_quality(document: Document) -> None:
    frame = pd.read_csv(TABLES / "table_00_data_quality.csv")
    frame = frame[frame["role"].isin(["main_experiment", "robustness_only"])]
    batch_labels = {"without_fails_1": "批次 1", "without_fails_2": "批次 2", "without_fails_3": "批次 3"}
    role_labels = {"main_experiment": "主实验", "robustness_only": "零样本稳健性"}
    rows = [
        [
            batch_labels[str(row.file_id)],
            role_labels[str(row.role)],
            f"{int(row.raw_rows):,}",
            f"{int(row.clean_rows):,}",
            f"{int(row.deleted_rows):,}",
            f"{100 * float(row.deletion_ratio):.2f}%",
        ]
        for row in frame.itertuples()
    ]
    add_table(
        document,
        "表 1  数据清洗与实验角色",
        ["批次", "角色", "原始请求", "清洗后", "删除", "删除率"],
        rows,
        "BurstGPT v2.0；主实验使用批次 1/2，批次 3 仅作零样本稳健性；删除项为全字段精确重复。",
        [2.3, 3.3, 2.5, 2.5, 2.1, 1.8],
    )


def insert_split_table(document: Document) -> None:
    frame = pd.read_csv(TABLES / "table_01_split_summary.csv")
    rows = [
        [
            str(row.split),
            f"{int(row.n_rows):,}",
            f"{100 * float(row.share_of_rows):.1f}%",
            f"{int(row.zero_windows):,}",
            f"{int(row.actual_burst_windows):,}",
            f"{100 * float(row.actual_burst_rate):.3f}%",
        ]
        for row in frame.itertuples()
    ]
    add_table(
        document,
        "表 2  主序列按目标时间的 70/15/15 切分",
        ["切分", "5 分钟窗", "占比", "零窗", "P95 突发窗", "突发率"],
        rows,
        "BurstGPT 批次 1/2；单位为窗口数；P95=315,246.2 Token/5min，仅由训练集确定。",
        [2.2, 2.5, 2.0, 2.2, 2.6, 2.3],
    )


def insert_model_table(document: Document) -> None:
    rows = [
        ["Persistence", "最后完整窗口 y(t−5min)", "无拟合"],
        ["Seasonal Naive", "目标时刻前 24h 窗口", "无拟合"],
        ["XGBoost", "36 个历史滞后、滚动、构成和相对相位特征", "12 候选；验证 MAE"],
        ["LSTM", "144×10 序列；输入结束于 t−5min", "32 单元；验证 MAE/epoch"],
    ]
    add_table(
        document,
        "表 3  模型输入与选择边界",
        ["模型", "在预测原点可用的输入", "选择/拟合边界"],
        rows,
        "所有任务目标为 t+h；scaler、训练 P95 和拟合统计量只使用训练范围，测试集只在冻结后报告。",
        [3.0, 7.4, 4.7],
    )


def insert_main_results(document: Document) -> None:
    add_table(
        document,
        "表 4  BurstGPT 批次 1/2 的 15 分钟主测试总体结果",
        ["模型", "MAE", "RMSE", "sMAPE", "较季节基线", "AP"],
        exact_main_rows(),
        "n=5,228；MAE/RMSE 单位为 Token/5min；AP 为非插值 average precision。模型按验证证据冻结。",
        [2.7, 2.3, 2.3, 2.3, 2.6, 2.3],
    )
    add_figure(document, FIGURES / "fig_13_model_mae_f1_h15.png", "图 13  15 分钟总体 MAE 与固定阈值 F1（BurstGPT 批次 1/2 测试集）")
    add_figure(document, FIGURES / "fig_11_test_week_forecasts_h15.png", "图 11  15 分钟预测在客观选择代表周的曲线对比（Token/5min）")


def insert_burst_results(document: Document) -> None:
    frame = pd.read_csv(TABLES / "table_05_burst_results.csv")
    rows = [
        [
            str(row.model),
            f"{float(row.precision):.3f}",
            f"{float(row.recall):.3f}",
            f"{float(row.f1):.0f}",
            f"{float(row.pr_auc):.4f}",
            f"{int(row.true_positive)}/{int(row.false_positive)}/{int(row.false_negative)}/{int(row.true_negative)}",
        ]
        for row in frame.itertuples()
    ]
    add_table(
        document,
        "表 5  BurstGPT 批次 1/2 的 15 分钟固定 P95 突发结果",
        ["模型", "P", "R", "F1", "AP", "TP/FP/FN/TN"],
        rows,
        "n=5,228，真实突发窗=5；P/R 分别为 Precision/Recall；阈值为训练集 P95=315,246.2 Token/5min，未使用测试标签调整。",
        [2.7, 2.3, 2.0, 1.7, 2.4, 3.5],
    )
    add_figure(document, FIGURES / "fig_12_burst_window_zoom_h15.png", "图 12  最大真实负载附近预测与训练 P95（Token/5min）")
    add_figure(document, FIGURES / "fig_14_pr_curves_h15.png", "图 14  15 分钟突发排序的 Precision–Recall 曲线")


def insert_segment_results(document: Document) -> None:
    document.add_page_break()
    frame = pd.read_csv(TABLES / "table_06_segment_errors.csv")
    frame = frame[(frame["dimension"] == "actual_burst")]
    rows = [
        [
            str(r.model),
            "突发" if str(r.segment) == "Burst" else "非突发",
            f"{int(r.n_observations):,}",
            f"{float(r.mae):,.1f}",
            f"{float(r.mean_bias_y_pred_minus_y_true):,.1f}",
        ]
        for r in frame.itertuples()
    ]
    add_table(
        document,
        "表 6  BurstGPT 批次 1/2 的 15 分钟突发/非突发分组误差",
        ["模型", "分组", "n", "MAE", "平均偏差"],
        rows,
        "MAE 和平均偏差单位为 Token/5min；偏差=预测−真实；分组只用于事后诊断，不进入训练或模型选择。",
        [3.2, 2.4, 1.8, 3.1, 3.1],
    )
    add_figure(document, FIGURES / "fig_16_segment_error_burst_h15.png", "图 16  真实突发与非突发窗口的 15 分钟 MAE（Token/5min）")


def insert_ablation(document: Document) -> None:
    document.add_page_break()
    frame = pd.read_csv(TABLES / "table_06a_xgb_h15_ablation.csv")
    labels = {
        "full_features": "完整特征",
        "without_service_structure": "去服务构成",
        "lags_plus_relative_calendar": "滞后+相对相位",
    }
    rows = [
        [
            labels[str(r.feature_set)],
            str(r.split),
            str(int(r.n_features)),
            f"{float(r.mae):,.0f}",
            str(int(r.validation_mae_rank)),
            "是" if bool(r.selected_on_validation) else "否",
        ]
        for r in frame.itertuples()
    ]
    add_table(
        document,
        "表 6a  XGBoost 15 分钟最小特征消融",
        ["特征组", "切分", "特征数", "MAE", "验证名次", "入选"],
        rows,
        "BurstGPT 批次 1/2；MAE 单位 Token/5min。测试结果不得回写为新选择。",
        [4.2, 1.8, 1.5, 2.6, 2.0, 2.0],
    )


def insert_robustness(document: Document) -> None:
    document.add_page_break()
    frame = pd.read_csv(TABLES / "table_07_robustness_burstgpt3.csv")
    frame = frame[frame["horizon_minutes"] == 15]
    rows = [
        [
            str(r.model),
            f"{float(r.mae):,.0f}",
            f"{float(r.rmse):,.0f}",
            f"{float(r.smape_percent):.1f}%",
            f"{float(r.f1):.3f}",
            f"{float(r.pr_auc):.3f}",
        ]
        for r in frame.itertuples()
    ]
    add_table(
        document,
        "表 7  BurstGPT 第 3 批的 15 分钟零样本跨时期结果",
        ["模型", "MAE", "RMSE", "sMAPE", "F1", "AP"],
        rows,
        "n=31,389；MAE/RMSE 单位 Token/5min；模型、scaler 和 P95 全部来自批次 1/2，批次 3 拟合/选择行数为 0。",
        [2.7, 2.5, 2.5, 2.2, 2.1, 2.1],
    )
    robust_fig = FIGURES / "w5_python" / "fig_w5_py_10_burstgpt3_robustness.png"
    if robust_fig.exists():
        add_figure(document, robust_fig, "图 19  BurstGPT 第 3 批零样本 5/15/60 分钟稳健性：总体 MAE 与突发 F1")


REFERENCES = [
    "[1] Wang Y, Chen Y, Li Z, et al. BurstGPT: A Real-World Workload Dataset to Optimize LLM Serving Systems. KDD, 2025. doi:10.1145/3711896.3737413.",
    "[2] Cortez E, Bonde A, Muzio A, et al. Resource Central: Understanding and Predicting Workloads for Improved Resource Management in Large Cloud Platforms. SOSP, 2017:153–167. doi:10.1145/3132747.3132772.",
    "[3] Shahrad M, Fonseca R, Goiri I, et al. Serverless in the Wild: Characterizing and Optimizing the Serverless Workload at a Large Cloud Provider. USENIX ATC, 2020:205–218.",
    "[4] Yu G I, Jeong J S, Kim G W, et al. Orca: A Distributed Serving System for Transformer-Based Generative Models. OSDI, 2022:521–538.",
    "[5] Kwon W, Li Z, Zhuang S, et al. Efficient Memory Management for Large Language Model Serving with PagedAttention. SOSP, 2023:611–626. doi:10.1145/3600006.3613165.",
    "[6] Li Z, Zheng L, Zhong Y, et al. AlpaServe: Statistical Multiplexing with Model Parallelism for Deep Learning Serving. OSDI, 2023:663–679.",
    "[7] Patel P, Choukse E, Zhang C, et al. Splitwise: Efficient Generative LLM Inference Using Phase Splitting. ISCA, 2024:118–132. doi:10.1109/ISCA59077.2024.00019.",
    "[8] Zhong Y, Liu S, Chen J, et al. DistServe: Disaggregating Prefill and Decoding for Goodput-optimized Large Language Model Serving. OSDI, 2024:193–210.",
    "[9] Stojkovic J, Zhang C, Goiri I, et al. DynamoLLM: Designing LLM Inference Clusters for Performance and Energy Efficiency. HPCA, 2025. doi:10.1109/HPCA61900.2025.00102.",
    "[10] Chen T, Guestrin C. XGBoost: A Scalable Tree Boosting System. KDD, 2016:785–794. doi:10.1145/2939672.2939785.",
    "[11] Hochreiter S, Schmidhuber J. Long Short-Term Memory. Neural Computation, 1997, 9(8):1735–1780. doi:10.1162/neco.1997.9.8.1735.",
    "[12] Makridakis S, Spiliotis E, Assimakopoulos V. The M4 Competition: 100,000 Time Series and 61 Forecasting Methods. International Journal of Forecasting, 2020, 36(1):54–74. doi:10.1016/j.ijforecast.2019.04.014.",
]


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    for style_name, size, color in (("Title", 24, NAVY), ("Heading 1", 16, NAVY), ("Heading 2", 12, BLUE)):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(6)

    caption = styles["Caption"]
    caption.font.name = "Microsoft YaHei"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("LLM 服务 Token 负载预测与突发流量识别 · 论文初稿")
    set_run_font(run, size=8.5, color=MID_GRAY)
    add_page_number(section.footer.paragraphs[0])


def add_cover(document: Document) -> None:
    document.add_paragraph()
    document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(48)
    title.paragraph_format.space_after = Pt(24)
    run = title.add_run("LLM 服务 Token 负载预测")
    set_run_font(run, size=25, bold=True, color=NAVY)
    run.add_break()
    run = title.add_run("与突发流量识别")
    set_run_font(run, size=25, bold=True, color=NAVY)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("——基于 BurstGPT 真实轨迹的 XGBoost 与 LSTM 比较")
    set_run_font(run, size=16, bold=True, color=BLUE)
    rule = document.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = rule.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    set_run_font(run, size=10, color=BLUE)
    badge = document.add_paragraph()
    badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    badge.paragraph_format.space_before = Pt(22)
    run = badge.add_run("可复现研究包 · 论文初稿 · 2026-07-21")
    set_run_font(run, size=11, bold=True, color=NAVY)
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(16)
    run = meta.add_run("主任务：未来 15 分钟  |  粒度：Token/5min  |  Seed：42")
    set_run_font(run, size=10, color=MID_GRAY)
    document.add_page_break()


def build_from_markdown(document: Document) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    paragraph_buffer: list[str] = []
    skip_references_tail = False

    def flush() -> None:
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        text = " ".join(line.strip() for line in paragraph_buffer).strip()
        paragraph_buffer = []
        if not text:
            return
        paragraph = document.add_paragraph()
        if text.startswith(("N_k =", "y_k =")):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            run = paragraph.add_run(text.strip("`"))
            set_run_font(run, "Cambria Math", 10.5, color=NAVY)
        else:
            add_inline_markup(paragraph, text)

        if text.startswith("聚合后补齐完整规则网格"):
            insert_data_quality(document)
        elif text.startswith("公开时间为相对秒"):
            insert_split_table(document)
        elif text.startswith("时间戳表示窗口左边界"):
            insert_model_table(document)
        elif text.startswith("**观察到什么。** LSTM"):
            insert_main_results(document)
        elif text.startswith("**观察到什么。** 四种方法"):
            insert_burst_results(document)
        elif text.startswith("**观察到什么。** 突发窗误差"):
            insert_segment_results(document)
        elif text.startswith("**观察到什么。** 完整 36 特征"):
            insert_ablation(document)
        elif text.startswith("**观察到什么。** 跨时期模型"):
            insert_robustness(document)

    for line in lines:
        if line.startswith("# "):
            continue
        if line == "## 参考文献":
            flush()
            document.add_heading("参考文献", level=1)
            for item in REFERENCES:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.first_line_indent = Cm(-0.74)
                paragraph.paragraph_format.left_indent = Cm(0.74)
                add_inline_markup(paragraph, item)
            skip_references_tail = True
            continue
        if skip_references_tail:
            continue
        if line.startswith("## "):
            flush()
            document.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            flush()
            document.add_heading(line[4:].strip(), level=2)
        elif not line.strip():
            flush()
        else:
            paragraph_buffer.append(line)
    flush()


def main() -> None:
    if not (ROOT / "artifacts" / "final" / "manifest_sha256.csv").exists():
        raise FileNotFoundError("Run scripts/package_final_artifacts.py first")
    DELIVERABLES.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)
    add_cover(document)
    build_from_markdown(document)

    properties = document.core_properties
    properties.title = "LLM 服务 Token 负载预测与突发流量识别"
    properties.subject = "BurstGPT, XGBoost, LSTM, reproducible workload forecasting"
    properties.author = "Research package"
    properties.keywords = "LLM serving; token load; forecasting; burst; XGBoost; LSTM"
    document.save(OUTPUT)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    main()
